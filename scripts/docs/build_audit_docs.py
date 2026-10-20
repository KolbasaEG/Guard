"""Create the approved Word manuals. Uses python-docx from the bundled runtime."""
from __future__ import annotations
import hashlib, json, re, subprocess
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image
from audit_content import TECH, SCENARIOS, EVENT_TEXT

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/audit-package';WORK=ROOT/'Guard/obj/audit-docs';SHOTS=WORK/'screenshots'
TEAL='264B59';GRAY='54616B'
GROUPS=['Первый запуск','Контроль и поиск событий','Карточка события и срок хранения',
 'Пределы очереди и включение SIEM','Локальный режим и возобновление очереди',
 'Перенос очереди и активная отправка','Сертификаты и доступность сервера',
 'Ошибки TLS и поиск в RuSIEM','Рост очереди и превышение возраста',
 'Переполнение и отказ локальной записи','Identity и отказ входа Root',
 'Ограниченная сессия и возврат к работе','Обычная сессия и отзыв полномочий',
 'Версия настроек и неизвестный commit','Выход и ошибки очистки',
 'Дисковое место и перезапуск','Обновление и восстановление копии',
 'Экспорт и обращение в сопровождение','Полнота событий и повторы',
 'Неизвестные показатели и технические логи','Тест соединения и ввод настроек']

def set_font(run,size=None,bold=None,color=None):
    run.font.name='Calibri';rf=run._element.get_or_add_rPr();fonts=rf.rFonts
    if fonts is None:fonts=OxmlElement('w:rFonts');rf.insert(0,fonts)
    for key in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+key),'Calibri')
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if color:run.font.color.rgb=RGBColor.from_string(color)
    return run

class Manual:
    def __init__(self,title,subtitle):
        self.doc=Document();self.entries=[];self.page=0;self.bookmark=0
        s=self.doc.sections[0];s.page_width=Cm(21.59);s.page_height=Cm(27.94)
        s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.5);s.left_margin=Cm(2);s.right_margin=Cm(2)
        s.header_distance=Cm(.7);s.footer_distance=Cm(.7)
        styles=self.doc.styles
        for name in ('Normal','Body Text','List Number','List Bullet'):
            st=styles[name];st.font.name='Calibri';st.font.size=Pt(11)
            st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.06
        for name,size,col in [('Title',30,'000000'),('Subtitle',15,'000000'),('Heading 1',18,'000000'),('Heading 2',12,'000000'),('Heading 3',11,'000000')]:
            st=styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(col)
            st.paragraph_format.space_after=Pt(9);st.paragraph_format.space_before=Pt(8)
            st.paragraph_format.keep_with_next=True
        styles['Title'].font.bold=False
        for border in list(styles['Title'].element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        self.doc.core_properties.author='Guard';self.doc.core_properties.title=title
        self.doc.core_properties.subject=subtitle;self.doc.core_properties.keywords='Guard, аудит, Root, RuSIEM, BPMN'
        settings=self.doc.settings.element;up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');settings.append(up)
        lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');styles['Normal'].element.get_or_add_rPr().append(lang)
        foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        set_font(foot.add_run('Guard  |  Аудит  |  Редакция 1.0  |  '),8,color=GRAY)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
        self.newpage();self.doc.add_paragraph('GUARD',style='Subtitle')
        self.doc.add_paragraph(title,style='Title');self.doc.add_paragraph(subtitle,style='Subtitle')
        self.p('Редакция 1.0   20 октября 2026 года')
        self.p('Назначение документа',bold=True)
        self.p('Документ описывает действующий механизм аудита и порядок его проверки перед эксплуатацией. Локальная запись обязательна в обоих режимах. При неисправности обычная работа запрещается; Root использует отдельный ограниченный путь диагностики и восстановления.')
        self.table(['Сведения','Значение'],[['Продукт','Guard'],['Снимок реализации','codex/logging-reliability от 20 октября 2026 года'],['Аудитория','Root, сопровождение и ответственные за информационную безопасность'],['Комплект','Техническое описание, инструкция Root и модель BPMN 2.0'],['Эксплуатационная приёмка','Реальная RuSIEM, нагрузка и восстановление принимаются отдельно']],[4.2,12.8])
        self.p('Оглавление содержит переходы к разделам. Для подробного поиска сценария используйте коды R01–R42 и область навигации Word. Коды P00–P10 относятся к схемам BPMN.')
        self.newpage();self.doc.add_paragraph('Оглавление',style='Heading 1');self.toc_anchor=self.doc.paragraphs[-1]
        self.toc_paragraphs=[]
    def newpage(self):
        if self.page:self.doc.add_page_break()
        self.page+=1
    def heading(self,title,level=1,key=None,toc=True):
        p=self.doc.add_paragraph(title,style='Heading '+str(level))
        self.bookmark+=1;name=key or 'section_'+str(self.bookmark)
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(self.bookmark));start.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(self.bookmark));p._p.insert(0,start);p._p.append(end)
        if toc:self.entries.append((title,name,self.page))
        return p
    def p(self,text,bold=False,size=None):
        p=self.doc.add_paragraph();set_font(p.add_run(text),size,bold)
        return p
    def table(self,headers,rows,widths=None):
        table=self.doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        table.style='Table Grid'
        borders=OxmlElement('w:tblBorders')
        for side in ('top','left','bottom','right','insideH','insideV'):
            border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        table._tbl.tblPr.append(borders)
        if widths is None:widths=[17/len(headers)]*len(headers)
        for c,width in zip(table.columns,widths):c.width=Cm(width)
        for i,h in enumerate(headers):table.rows[0].cells[i].text=h
        for row in rows:
            cells=table.add_row().cells
            for i,text in enumerate(row):cells[i].text=str(text)
        for rowidx,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
            if rowidx==0:
                repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for i,cell in enumerate(row.cells):
                cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                for side in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');m.append(el)
                tcpr.append(m)
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),TEAL if rowidx==0 else ('F3F6F7' if rowidx%2 else 'FFFFFF'));tcpr.append(sh)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.02
                    for run in p.runs:set_font(run,9.5,rowidx==0,'FFFFFF' if rowidx==0 else '202C33')
        self.doc.add_paragraph().paragraph_format.space_after=Pt(1)
        return table
    def steps(self,items):
        for i,item in enumerate(items,1):
            p=self.doc.add_paragraph();p.paragraph_format.left_indent=Cm(.3);p.paragraph_format.first_line_indent=Cm(-.3)
            set_font(p.add_run(str(i)+'  '),11,True);set_font(p.add_run(item),11)
    def blocks(self,blocks):
        for kind,data in blocks:
            if kind=='p':self.p(data)
            elif kind=='h':self.heading(data,2,toc=False)
            elif kind=='table':
                if len(data)==2 and data[0][0]=='Код':self.table(*data,[1.5,15.5] if data[1][0][0]=='L01' else [5.5,11.5])
                else:self.table(*data)
            elif kind=='steps':self.steps(data)
            elif kind=='note':self.p(data,bold=True)
    def toc(self):
        last=self.toc_anchor._p
        for title,key,page in self.entries:
            p=OxmlElement('w:p');pr=OxmlElement('w:pPr');space=OxmlElement('w:spacing');space.set(qn('w:after'),'110');pr.append(space);p.append(pr)
            link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),key);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),TEAL);rp.append(color);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');rp.append(sz);r.append(rp);txt=OxmlElement('w:t');txt.text=title+'  ·  '+str(page);r.append(txt);link.append(r);p.append(link);last.addnext(p);last=p
    def image(self,path,caption,width=17):
        with Image.open(path) as im:actual_width=min(width,21.5*im.width/im.height)
        self.doc.add_picture(str(path),width=Cm(actual_width));self.p(caption,size=9)
    def save(self,name):
        self.toc();path=OUT/name;self.doc.save(path)
        return {'file':name,'plannedPages':self.page,'bookmarks':len(self.entries),'toc':self.entries}

def enum_catalog():
    text=(ROOT/'Guard.Domain/Core/0-Enums/AuditEventType.cs').read_text(encoding='utf-8-sig')
    events=re.findall(r'^\s*(\w+)\s*=\s*(\d+)',text,re.M)
    refs={name:[] for name,_ in events}
    for base in ['Guard.Application','Guard.Infrastructure','Guard']:
        for path in (ROOT/base).rglob('*'):
            if not path.is_file() or path.suffix not in ('.cs','.razor') or any(x in path.parts for x in ('bin','obj','Migrations')):continue
            content=path.read_text(encoding='utf-8-sig',errors='replace')
            if path.name=='CefFormatter.cs':continue
            for name in re.findall(r'\bAuditEventType\.(\w+)',content):
                if name in refs:refs[name].append(path.relative_to(ROOT).as_posix())
    return events,refs

def technical(events,refs):
    m=Manual('Подсистема аудита Guard','Техническое описание')
    for title,blocks in TECH:
        m.newpage();m.heading(title);m.blocks(blocks)
    for start in range(0,len(events),12):
        subset=events[start:start+12];m.newpage();m.heading('Справочник событий '+str(start+1)+' '+str(start+len(subset)))
        rows=[]
        for name,value in subset:
            note=EVENT_TEXT[name]
            if not refs[name] and not note.startswith('Совместимое'):note+='; активный источник в снимке не найден'
            rows.append([value,name,note])
        m.table(['Код','EventType','Назначение и статус'],rows,[1.1,6.2,9.7])
        m.p('Числа сохраняются в БД и не перенумеровываются. Справочник перечисляет enum полностью; совместимые значения отдельно от действующих функций. Источники упоминаний перечислены в source-manifest.json.',size=9)
    m.newpage();m.heading('Термины и проверка результата')
    m.table(['Термин','Значение'],[['Root','Аккаунт с действующим назначением роли ROOT, проверенным через Identity'],['Pending','Обязательное событие без SentAtUtc'],['Local','Режим без нового обязательства отправки SIEM'],['Recovery only','Ограниченная сессия на 15 минут для диагностики'],['Commit','Подтверждённая фиксация транзакции'],['Unknown','Подтверждение результата потеряно; требуется проверка'],['EventId','Идентификатор неизменного события и ключ сопоставления повторов'],['CorrelationId','Связь нескольких событий одной операции'],['PayloadBytes','Размер данных события UTF-8, не физический размер БД'],['Lease','Временная аренда отправки на 2 минуты'],['CEF','Формат полей события для получателя'],['BPMN DI','Координаты и соединения для графического импорта модели']],[3.7,13.3])
    m.p('Критерий результата определяется границей операции: локальный commit, завершённая TLS запись, сохранение события в RuSIEM и фактическая передача файла — разные подтверждения. При неизвестном исходе проверяют факты до повторения.')
    return m.save('Guard_Audit_Technical.docx')

def root_guide():
    m=Manual('Управление аудитом Guard','Инструкция Root по настройке контролю и восстановлению')
    m.newpage();m.heading('Быстрый старт и порядок действий')
    m.p('Начните с определения режима и состояния. Наличие доступа Root к диагностике не разрешает выполнять обычные операции при неисправном аудите.')
    m.table(['Ситуация','Что делать'],[['Первый запуск NotConfigured','R01 для Local или R02 для RuSIEM'],['Очередь в пределах но растёт','R17 затем R14 R15 R16'],['Превышен возраст','R18'],['Превышены байты или количество','R19'],['Нет локальной записи','R20'],['Не подтверждается Root или Identity','R21 R22'],['После ремонта осталась recovery сессия','R24'],['Неизвестно сохранилось ли изменение','R28'],['Нужен безопасный выход','R29']],[7,10])
    m.heading('Безопасная последовательность восстановления',2,toc=False)
    m.steps(['Определите код и свежесть состояния.','Устраните причину в пределах доступных средств Root либо привлеките сопровождение.','Подтвердите свежую успешную проверку и состояние очереди.','Проверьте фактический результат прерванной операции.','После recovery выполните выход и обычный вход; при необходимости повторите исходное действие.'])
    m.p('Время состояния и карточки — UTC; время списка журнала — зона браузера. Для обращения в сопровождение всегда указывайте UTC и EventId.')
    m.newpage();m.heading('Интерфейс и подтверждение изменений')
    m.table(['Элемент','Назначение'],[['Боковое меню','Состояние и настройки, Диагностический журнал, Обычный вход'],['Верхняя панель','Переключатель темы, кнопка выхода'],['Состояние','Действующие параметры и результат последней проверки'],['Форма настроек','Несохранённые значения; изменение селектора ещё не применяет режим'],['Сохранить настройки','Подтверждает изменение с обязательным аудитом'],['Обновить','Перечитывает фактическую конфигурацию и заменяет несохранённый ввод'],['Проверить соединение и отправить тест','Проверяет поля формы без их автоматического сохранения'],['Карточка события','Исходные значения, причины, связь событий и попытки доставки']],[6.5,10.5])
    m.p('Перед сменой назначения с pending требуется отдельное подтверждение переноса. При конфликте версии обновите форму. При потере подтверждения сохранения сначала установите фактическое состояние. Подробные действия приведены в сценариях ниже.')
    m.p('Стандартный срок хранения — 1 календарный год, допустимый диапазон 1–30. Срок pending не ограничивается удалением по общему сроку: обязательные неотправленные события сохраняются до доставки.')
    for i in range(0,len(SCENARIOS),2):
        m.newpage();m.heading(GROUPS[i//2])
        for sc in SCENARIOS[i:i+2]:
            id,title,symptom,steps,expected,verify,escalation=sc
            m.heading(id+' '+title,2,key=id,toc=False)
            m.p('Признак  '+symptom)
            m.steps(steps)
            p=m.doc.add_paragraph();set_font(p.add_run('Результат  '),10.5,True,TEAL);set_font(p.add_run(expected),10.5)
            p=m.doc.add_paragraph();set_font(p.add_run('Проверка  '),10.5,True,TEAL);set_font(p.add_run(verify),10.5)
            p=m.doc.add_paragraph();set_font(p.add_run('Если не помогло  '),10.5,True,TEAL);set_font(p.add_run(escalation),10.5)
    m.newpage();m.heading('Справочник состояний и параметров')
    m.table(['Состояние','Действие Root'],[['Ready','Контролировать свежесть и очередь'],['NotConfigured','Подтвердить режим R01 R02'],['InvalidConfiguration','Проверить параметры R42'],['StorageUnavailable','Восстановить локальную запись R20'],['QueueAgeExceeded','Устранить старую очередь R18'],['QueueVolumeExceeded','Устранить превышение R19'],['ServiceReserveExceeded','Исчерпан резерв отдельной служебной вставки; диагностика и восстановление'],['Исход неизвестен','Проверить результат R28']] ,[5.3,11.7])
    m.table(['Настройка','Допустимое значение'],[['Хранение','1–30 календарных лет'],['Возраст SIEM','1–525600 минут'],['Объём SIEM','Не меньше 1024 байт'],['Количество','От 0; 0 снимает только этот дополнительный предел'],['Порт','1–65535'],['Сервер','Имя или IP без схемы и пути'],['CA','Только публичный PEM до 16384 символов'],['Recovery','15 минут, без продления и постоянной cookie']],[5.3,11.7])
    if SHOTS.exists():
        for filename,heading,caption in [
            ('settings.jpg','Страница состояния и настроек','Рисунок 1  Состояние и настройки на изолированном стенде. Значения примерные. Формат даты состояния — UTC.'),
            ('siem.jpg','Поля подключения RuSIEM','Рисунок 2  Два фрагмента формы подключения и обязательной очереди. Значения не применяются до сохранения.'),
            ('journal.jpg','Диагностический журнал','Рисунок 3  Поиск и карточка события доступны через диагностический журнал. Время списка соответствует зоне браузера.')]:
            path=SHOTS/filename
            if path.exists():m.newpage();m.heading(heading);m.image(path,caption)
    m.newpage();m.heading('Памятка и комплект документов')
    m.p('При блокировке сначала определите код. Не удаляйте pending, не изменяйте время постановки и не отключайте TLS проверку ради возвращения зелёного состояния. Root может согласованно изменить режим или пределы только с исправной обязательной локальной записью.')
    m.table(['Для сопровождения','Что записать'],[['Состояние','Код, режим, время UTC, свежесть, пределы и очередь'],['Операция','Действие, объект, ожидаемый и фактический результат'],['Связь','EventId, CorrelationId, ReasonCode и Outcome'],['Доставка','Адресат, времена, тип ошибки, следующая попытка'],['Контекст','Последние изменения, перезапуск, версия приложения'],['Исключить','Пароли, cookie, токены, приватные ключи, credentials подключений']],[4.5,12.5])
    m.p('Технические гарантии и ограничения L01–L11 раскрыты в Guard_Audit_Technical.docx. Схемы P00–P10 доступны в BPMN/Guard_Audit_Master.bpmn и BPMN/Guard_Audit_BPMN_Atlas.pdf. README.md описывает состав и импорт. Документы относятся к текущему снимку исходников; фактический выпуск и реальная RuSIEM принимаются отдельно.')
    return m.save('Guard_Audit_Root_Guide.docx')

def manifest(events,refs):
    paths=[ROOT/'docs/audit-compliance-plan.md',ROOT/'docs/audit-operations.md',ROOT/'Guard.Domain/Core/0-Enums/AuditEventType.cs']
    for folder in ['Guard.Infrastructure/Core/9-Logging','Guard.Infrastructure/Core/Identity','Guard/Components/Pages/Audit','Guard/Web/Identity']:
        paths.extend(p for p in (ROOT/folder).glob('*') if p.suffix in ('.cs','.razor','.css'))
    for file in ['Guard.Infrastructure/Core/6-Services/AuditAdministrationService.cs','Guard.Infrastructure/Core/6-Services/AuditSearchService.cs','Guard.Application/Core/6-Services/PermissionService.cs']:
        paths.append(ROOT/file)
    data={'date':'2026-10-20','branch':'codex/logging-reliability','workingTree':'Uncommitted implementation snapshot',
      'files':[{'path':p.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(set(paths))],
      'events':[{'code':int(code),'name':name,'purpose':EVENT_TEXT[name],'references':sorted(set(refs[name]))} for name,code in events]}
    (OUT/'source-manifest.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

def main():
    OUT.mkdir(parents=True,exist_ok=True);events,refs=enum_catalog();manifest(events,refs)
    result=[technical(events,refs),root_guide()]
    (WORK/'doc-structure.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':main()
